# matchwise_app.py – Professional Version (Updated with smaller, elegant gray graph)

import streamlit as st
import pandas as pd
import tempfile
import os
from datetime import datetime
from pdfminer.high_level import extract_text
import docx
from sentence_transformers import SentenceTransformer, util
from sklearn.feature_extraction.text import CountVectorizer
import numpy as np
import matplotlib.pyplot as plt
from fpdf import FPDF

# === Suppress warnings and logs ===
import warnings
import logging
warnings.filterwarnings("ignore")
logging.getLogger("pdfminer").setLevel(logging.ERROR)

# === Load Model ===
model = SentenceTransformer("multi-qa-MiniLM-L6-cos-v1")

# === Helper Functions ===
def extract_text_from_pdf(file):
    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
        tmp.write(file.read())
        return extract_text(tmp.name).strip()

def extract_text_from_docx(file):
    doc = docx.Document(file)
    return "\n".join([para.text for para in doc.paragraphs]).strip()

def extract_resume_text(uploaded_file):
    if uploaded_file.name.endswith(".pdf"):
        return extract_text_from_pdf(uploaded_file)
    elif uploaded_file.name.endswith(".docx"):
        return extract_text_from_docx(uploaded_file)
    return ""

def preprocess_text(text):
    return " ".join([line for line in text.lower().split("\n") if len(line.strip()) > 5])

def extract_keywords(text, top_n=15):
    vectorizer = CountVectorizer(stop_words='english', max_features=1000)
    vectors = vectorizer.fit_transform([text])
    keywords = vectorizer.get_feature_names_out()
    word_freq = zip(keywords, vectors.toarray()[0])
    sorted_keywords = sorted(word_freq, key=lambda x: x[1], reverse=True)
    return [kw for kw, _ in sorted_keywords[:top_n]]

def keyword_overlap_score(resume, jd):
    r_words = set(resume.split())
    j_words = set(jd.split())
    common = r_words & j_words
    return round(len(common) / len(j_words) * 100, 2) if j_words else 0.0

def generate_pdf_report(df):
    now = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt="MatchWise - Resume Compatibility Report", ln=True, align='C')
    pdf.cell(200, 10, txt=f"Generated: {now}", ln=True, align='C')
    pdf.ln(10)

    for index, row in df.iterrows():
        pdf.set_font("Arial", 'B', 11)
        pdf.cell(200, 10, txt=f"{index+1}. {row['Resume']}", ln=True)
        pdf.set_font("Arial", '', 11)
        pdf.cell(200, 10, txt=f" - Final Score: {row['Final Score %']}%", ln=True)
        pdf.cell(200, 10, txt=f" - BERT Match: {row['BERT Match %']}%, Keyword Match: {row['Keyword Match %']}%", ln=True)
        pdf.ln(5)

    output_path = f"matchwise_report_{now}.pdf"
    pdf.output(output_path)
    return output_path

# === Streamlit App Layout ===
st.set_page_config(page_title="MatchWise", layout="wide")
st.markdown("""
    <div style='text-align:center; padding: 10px 0;'>
        <h1 style='font-size: 36px;'>MatchWise</h1>
        <p style='font-size: 16px; color: gray;'>AI-Powered Resume–Job Compatibility Scoring</p>
    </div>
""", unsafe_allow_html=True)

col1, col2 = st.columns([1, 2])

# === JD Input and Resume Upload ===
with col1:
    st.subheader("Job Description Input")
    jd_input = st.text_area("Paste the job description:", height=300)

    st.subheader("Uploaded Resumes")
    uploaded_resumes = st.file_uploader("Upload multiple resumes (PDF or DOCX)", type=["pdf", "docx"], accept_multiple_files=True)

    if jd_input:
        jd_clean = preprocess_text(jd_input)
        jd_keywords = extract_keywords(jd_clean)
        st.markdown("**Top JD Keywords:**")
        st.code(", ".join(jd_keywords))

with col2:
    if st.button("Run Matching") and jd_input and uploaded_resumes:
        with st.spinner("Processing resumes..."):
            jd_embedding = model.encode(jd_clean, convert_to_tensor=True)
            results = []
            keyword_data = []

            for file in uploaded_resumes:
                resume_text = extract_resume_text(file)
                resume_clean = preprocess_text(resume_text)
                resume_keywords = extract_keywords(resume_clean)
                keyword_data.append((file.name, resume_keywords))

                # Match scoring
                resume_embedding = model.encode(resume_clean, convert_to_tensor=True)
                bert_score = util.pytorch_cos_sim(resume_embedding, jd_embedding).item()
                bert_percent = round(bert_score * 100, 2)
                keyword_percent = keyword_overlap_score(resume_clean, jd_clean)
                final_score = round(0.6 * bert_percent + 0.4 * keyword_percent, 2)

                results.append({
                    "Resume": file.name,
                    "BERT Match %": bert_percent,
                    "Keyword Match %": keyword_percent,
                    "Final Score %": final_score
                })

            df = pd.DataFrame(results).sort_values("Final Score %", ascending=False).reset_index(drop=True)

            # === Display Graph ===
            st.subheader("Match Score Overview")
            fig, ax = plt.subplots(figsize=(6, 4))  # smaller size
            ax.barh(df['Resume'], df['Final Score %'], color="gray")
            ax.invert_yaxis()
            ax.set_xlabel("Final Score %")
            ax.set_title("Resume Match Scores")
            st.pyplot(fig)

            # === Display Table ===
            st.subheader("Detailed Match Scores")
            st.dataframe(df, use_container_width=True)

            # === Keyword Comparison ===
            st.subheader("Resume Keyword Highlights")
            for name, keywords in keyword_data:
                st.markdown(f"**{name}**")
                st.code(", ".join(keywords))

            # === PDF Download ===
            pdf_file_path = generate_pdf_report(df)
            with open(pdf_file_path, "rb") as f:
                st.download_button(
                    label="Download PDF Report",
                    data=f,
                    file_name=os.path.basename(pdf_file_path),
                    mime="application/pdf"
                )
